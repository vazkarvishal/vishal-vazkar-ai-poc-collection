"""Deterministic evaluation and delivery for normalized UK property candidates."""
from __future__ import annotations

import argparse
import json
import os
import re
import smtplib
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit

from openpyxl import Workbook, load_workbook

DEFAULT_CONFIG = {"areas": [], "max_price": 800000, "min_bedrooms": 4, "min_bathrooms": 3, "min_sqft": 1500, "min_qualifying_bedrooms": 2, "min_bedroom_length_m": 4, "min_bedroom_width_m": 3, "max_station_minutes": 20, "max_outliers": 5, "name": "UK Property Watch"}
WEIGHTS = {"location": 15, "station": 10, "school": 10, "flood": 10, "environment": 5, "size_layout": 15, "value": 15, "kitchen": 8, "garage_garden": 7, "condition": 3, "amenities": 2}
HEADERS = ["Run date/time", "Property Key", "Area Priority", "Area", "Address/listing title", "Asking/guide price", "Source/portal", "Listing URL", "Listing status", "Bedrooms", "Bathrooms/WCs", "Internal sq ft", "Largest bedroom dimensions", "Second-largest bedroom dimensions", "Kitchen/layout note", "Garage", "Garden", "Parking", "Station used", "Estimated drive/taxi minutes", "Crime assessment", "School/catchment assessment", "Flood risk", "Road/noise/environmental note", "Calculated score", "Final status", "Primary rejection reason", "Secondary rejection reasons"]


def norm(value):
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", str(value or "").lower())).strip()


def key(p):
    address = norm(p.get("address"))
    postcode = norm(p.get("postcode"))
    if address and postcode:
        return f"address:{address}:{postcode}"
    url = p.get("url", "")
    if url:
        parts = urlsplit(url)
        canonical = urlunsplit((parts.scheme.lower(), parts.netloc.lower(), parts.path.rstrip("/"), "", ""))
        return f"listing:{canonical}:{norm(p.get('agent'))}"
    return f"fallback:{norm(p.get('title'))}:{norm(p.get('area'))}:{p.get('price')}:{norm(p.get('agent'))}"


def fingerprint(p):
    fields = ("price", "status", "floorplan_version", "auction_deadline", "listing_version", "bedrooms", "bathrooms", "sqft", "bedroom_dimensions")
    return json.dumps({f: p.get(f) for f in fields}, sort_keys=True)


def score(p):
    scores = p.get("scores") or {}
    if set(scores) != set(WEIGHTS):
        return None
    for name, weight in WEIGHTS.items():
        value = scores[name]
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not 0 <= value <= weight:
            raise ValueError(f"Invalid score {name}: expected 0..{weight}")
    return sum(scores.values())


def classify(p, config=None):
    config = DEFAULT_CONFIG | (config or {})
    reasons = []
    def fail(condition, reason):
        if condition:
            reasons.append(reason)
    fail(p.get("detached") is False, "NOT_DETACHED")
    fail(p.get("price") is not None and p["price"] > config["max_price"], "PRICE_OVER_MAX")
    fail(p.get("bedrooms") is not None and p["bedrooms"] < config["min_bedrooms"], "FEWER_THAN_MIN_BEDS")
    fail(p.get("bathrooms") is not None and p["bathrooms"] < config["min_bathrooms"], "FEWER_THAN_MIN_BATHS_WCS")
    fail(p.get("sqft") is not None and p["sqft"] < config["min_sqft"], "UNDER_MIN_SQFT")
    dimensions = p.get("bedroom_dimensions")
    if dimensions is not None:
        fail(sum(1 for pair in dimensions if len(pair) == 2 and max(pair) >= config["min_bedroom_length_m"] and min(pair) >= config["min_bedroom_width_m"]) < config["min_qualifying_bedrooms"], "BEDROOM_DIMENSIONS_FAIL")
    fail(p.get("garage") is False, "NO_GARAGE")
    fail(p.get("garden") is False, "NO_GARDEN")
    fail(p.get("large_kitchen") is False, "KITCHEN_LAYOUT_FAIL")
    fail(p.get("station_minutes") is not None and p["station_minutes"] > config["max_station_minutes"], "STATION_OVER_MAX_MINUTES")
    for field, code in (("crime_risk", "CRIME_RISK"), ("flood_risk", "FLOOD_RISK"), ("road_noise_risk", "ROAD_NOISE_RISK"), ("school_access_weak", "SCHOOL_ACCESS_WEAK")):
        fail(p.get(field) is True, code)
    if reasons:
        if p.get("outlier_reason") and p.get("detached") is True and not any(r in reasons for r in ("CRIME_RISK", "FLOOD_RISK", "ROAD_NOISE_RISK")):
            return "OUTLIER_CANDIDATE", reasons
        return "REJECTED", reasons
    required = ("detached", "price", "bedrooms", "bathrooms", "sqft", "bedroom_dimensions", "garage", "garden", "large_kitchen", "station_minutes", "crime_risk", "flood_risk", "road_noise_risk", "school_access_weak")
    missing = [f for f in required if p.get(f) is None]
    if missing or not p.get("url") or score(p) is None:
        return "UNASSESSED", ["LISTING_DATA_INSUFFICIENT", *missing]
    if score(p) < 65:
        return "REJECTED", ["SCORE_BELOW_65"]
    return "PRIMARY", []


def area_rank(p, config=None):
    areas = (config or DEFAULT_CONFIG)["areas"]
    try:
        return areas.index(p.get("area")) + 1
    except ValueError:
        return 99


def evaluate(candidates, history, config=None):
    config = DEFAULT_CONFIG | (config or {})
    unique = {}
    for p in candidates:
        if config["areas"] and p.get("area") not in config["areas"]:
            continue
        k = key(p)
        # Prefer the richer record when two portals expose the same address.
        if k not in unique or sum(v is not None for v in p.values()) > sum(v is not None for v in unique[k].values()):
            unique[k] = p
    rows = []
    for k, p in unique.items():
        state = "NEW" if k not in history else ("UNCHANGED" if fingerprint(p) == history[k]["fingerprint"] else "CHANGED")
        status, reasons = classify(p, config)
        rows.append({"key": k, "property": p, "state": state, "status": status, "reasons": reasons, "score": score(p)})
    outliers = sorted((r for r in rows if r["status"] == "OUTLIER_CANDIDATE"), key=lambda r: (area_rank(r["property"], config), -(r["score"] or 0), r["property"].get("price") or 0))
    for r in outliers[:config["max_outliers"]]:
        r["status"] = "OUTLIER"
    for r in outliers[config["max_outliers"]:]:
        r["status"] = "REJECTED"
        r["reasons"].insert(0, "OUTLIER_LIMIT_5")
    accepted = sorted((r for r in rows if r["status"] in ("PRIMARY", "OUTLIER") and r["state"] != "UNCHANGED"), key=lambda r: (area_rank(r["property"], config), -(r["score"] or 0), r["property"].get("price") or 0))
    rejected_new = sorted((r for r in rows if r["state"] == "NEW" and r["status"] not in ("PRIMARY", "OUTLIER")), key=lambda r: (area_rank(r["property"], config), r["status"], -(r["score"] or 0)))
    counts = {"skimmed": len(rows), "new": sum(r["state"] == "NEW" for r in rows), "unchanged": sum(r["state"] == "UNCHANGED" for r in rows), "changed": sum(r["state"] == "CHANGED" for r in rows), "primary": sum(r["status"] == "PRIMARY" for r in rows), "outliers": sum(r["status"] == "OUTLIER" for r in rows), "rejected_new": len(rejected_new)}
    counts["accepted"] = counts["primary"] + counts["outliers"]
    assert counts["skimmed"] == counts["new"] + counts["unchanged"] + counts["changed"]
    return rows, accepted, rejected_new, counts


def workbook(path, rows, now, config=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Newly reviewed rejected"
    ws.append(HEADERS)
    for r in rows:
        p = r["property"]
        dims = p.get("bedroom_dimensions") or []
        dimension = lambda i: " x ".join(map(str, dims[i])) if len(dims) > i else ""
        ws.append([now, r["key"], area_rank(p, config), p.get("area"), p.get("address") or p.get("title"), p.get("price"), p.get("source"), p.get("url"), p.get("status"), p.get("bedrooms"), p.get("bathrooms"), p.get("sqft"), dimension(0), dimension(1), p.get("kitchen_note"), p.get("garage"), p.get("garden"), p.get("parking"), p.get("station"), p.get("station_minutes"), p.get("crime_note"), p.get("school_note"), p.get("flood_note"), p.get("environment_note"), r["score"], r["status"], r["reasons"][0] if r["reasons"] else "", ", ".join(r["reasons"][1:])])
        if p.get("url"):
            cell = ws.cell(ws.max_row, 8)
            cell.hyperlink = p["url"]
            cell.style = "Hyperlink"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)
    check = load_workbook(path)
    for row in range(2, check.active.max_row + 1):
        c = check.active.cell(row, 8)
        if c.value and (not c.hyperlink or c.hyperlink.target != c.value):
            raise RuntimeError("Workbook hyperlink verification failed")


def render(accepted, counts, history_available, send_empty, attach):
    lines = []
    for r in accepted:
        p = r["property"]
        lines.extend([f"{p.get('address') or p.get('title')}, {p['area']} | £{p.get('price', 0):,} | {r['score'] if r['score'] is not None else 'unverified'}/100", p.get("url", "")])
        highlights = list(p.get("highlights", []))[:4 if r["status"] == "OUTLIER" else 5]
        if r["status"] == "OUTLIER":
            highlights.insert(0, "Why outlier: " + p["outlier_reason"])
        lines.extend("- " + h for h in highlights)
        lines.append("")
    if not accepted:
        lines.append("No new or materially changed matches.")
    lines += ["Technical run info", f"Unique properties skimmed this run: {counts['skimmed']}", f"Newly reviewed today: {counts['new']}", f"Previously reviewed unchanged: {counts['unchanged']}", f"Previously reviewed with material change: {counts['changed']}", f"Primary matches accepted: {counts['primary']}", f"Investment outliers accepted: {counts['outliers']}", f"Total accepted: {counts['accepted']}", f"Newly reviewed rejected / unassessed: {counts['rejected_new']}", f"Newly-reviewed rejection workbook: {'attached' if attach else 'not attached'}", f"Cross-run history status: {'available' if history_available else 'unavailable'}", f"SEND_EMPTY_EMAIL: {str(send_empty).lower()}"]
    return "\n".join(lines)


def send_email(subject, body, attachment):
    required = ("SMTP_HOST", "SMTP_USER", "SMTP_PASSWORD", "EMAIL_TO")
    if any(not os.getenv(v) for v in required):
        raise RuntimeError("Missing SMTP_HOST, SMTP_USER, SMTP_PASSWORD or EMAIL_TO")
    msg = EmailMessage()
    msg["From"] = os.getenv("EMAIL_FROM", os.environ["SMTP_USER"])
    msg["To"] = os.environ["EMAIL_TO"]
    msg["Subject"] = subject
    msg.set_content(body)
    if attachment:
        msg.add_attachment(attachment.read_bytes(), maintype="application", subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=attachment.name)
    with smtplib.SMTP_SSL(os.environ["SMTP_HOST"], int(os.getenv("SMTP_PORT", "465"))) as smtp:
        smtp.login(os.environ["SMTP_USER"], os.environ["SMTP_PASSWORD"])
        smtp.send_message(msg)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, help="JSON settings; see examples/reading-config.json")
    parser.add_argument("--input", type=Path, required=True, help="JSON array of normalized property candidates")
    parser.add_argument("--history", type=Path, default=Path("history.json"))
    parser.add_argument("--output", type=Path, default=Path("output"))
    parser.add_argument("--send-email", action="store_true")
    parser.add_argument("--send-empty-email", action=argparse.BooleanOptionalAction, default=True)
    parser.add_argument("--attach-rejected-xlsx", action=argparse.BooleanOptionalAction, default=True)
    args = parser.parse_args()
    config = DEFAULT_CONFIG | (json.loads(args.config.read_text()) if args.config else {})
    candidates = json.loads(args.input.read_text())
    if not isinstance(candidates, list):
        parser.error("Input must be a JSON array")
    history_available = args.history.exists()
    history = json.loads(args.history.read_text()) if history_available else {}
    rows, accepted, rejected, counts = evaluate(candidates, history, config)
    now = datetime.now().astimezone()
    args.output.mkdir(parents=True, exist_ok=True)
    attachment = args.output / f"Property_Watch_Newly_Reviewed_Rejected_{now.date()}.xlsx" if args.attach_rejected_xlsx else None
    if attachment:
        workbook(attachment, rejected, now.isoformat(timespec="seconds"), config)
    body = render(accepted, counts, history_available, args.send_empty_email, bool(attachment))
    (args.output / "report.txt").write_text(body + "\n")
    (args.output / "ledger.json").write_text(json.dumps([{k: v for k, v in r.items() if k != "property"} | {"property": r["property"]} for r in rows], indent=2))
    if args.send_email and (accepted or args.send_empty_email):
        send_email("🏠 " + config["name"] + " — " + ("New Matches" if accepted else "No New Matches"), body, attachment)
    # Commit state only after output generation and any requested email succeeds.
    for r in rows:
        p = r["property"]
        old = history.get(r["key"], {})
        history[r["key"]] = {"first_seen": old.get("first_seen", now.isoformat()), "last_seen": now.isoformat(), "fingerprint": fingerprint(p), "price": p.get("price"), "status": p.get("status"), "classification": r["status"]}
    temp = args.history.with_suffix(args.history.suffix + ".tmp")
    args.history.parent.mkdir(parents=True, exist_ok=True)
    temp.write_text(json.dumps(history, indent=2))
    temp.replace(args.history)
    print(body)


if __name__ == "__main__":
    main()
